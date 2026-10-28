"""
Excel styling engine using openpyxl.
Implements the 2-row merged header hierarchy, custom typography,
zebra-striping, and number formatting.
"""
import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.config import EXCEL_STYLES
from src.core.calculator import compute_profit_analysis_df


def style_profit_analysis_sheet(ws, df_profit):
    """
    Applies the custom 2-Row Header Layout to the 'Profit & Cost Analysis' sheet:
      Row 1: Main Categories (Merged headers for singles, merged E1:G1 for Return Breakup)
      Row 2: Sub-breakup headers for Returns & Cancellations
      Row 3+: Direct 1-to-1 data rows with alternating striping and numeric formatting.
    """
    ws.freeze_panes = "A3"

    header_fill = PatternFill(
        start_color=EXCEL_STYLES['HEADER_FILL_PRIMARY'],
        end_color=EXCEL_STYLES['HEADER_FILL_PRIMARY'],
        fill_type="solid"
    )
    header_fill_accent = PatternFill(
        start_color=EXCEL_STYLES['HEADER_FILL_ACCENT'],
        end_color=EXCEL_STYLES['HEADER_FILL_ACCENT'],
        fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color=EXCEL_STYLES['HEADER_FONT_COLOR'])
    header_font_sub = Font(name="Calibri", size=10, bold=True, color=EXCEL_STYLES['HEADER_SUB_COLOR'])
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color=EXCEL_STYLES['BORDER_COLOR']),
        right=Side(style='thin', color=EXCEL_STYLES['BORDER_COLOR']),
        top=Side(style='thin', color=EXCEL_STYLES['BORDER_COLOR']),
        bottom=Side(style='thin', color=EXCEL_STYLES['BORDER_COLOR'])
    )

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 24

    row1_cells = {
        'A1': 'SKU Name',
        'B1': 'Style name',
        'C1': 'Payout',
        'D1': 'Gross Units',
        'E1': 'Returned & Cancelled Units (Breakup)',
        'H1': 'Net Units',
        'I1': 'Net Production Cost [ Unit * ]',
        'J1': 'Bank Settlement [Projected] (INR)',
        'K1': 'PROFIT [ Bank Settelement - Net Production Cost ]'
    }
    for cell_ref, val in row1_cells.items():
        ws[cell_ref] = val

    ws['E2'] = 'RTO (Logistics Return)'
    ws['F2'] = 'RVP (Customer Return)'
    ws['G2'] = 'Cancelled Units'

    for col_idx in range(1, 12):
        col_letter = get_column_letter(col_idx)
        for r in [1, 2]:
            cell = ws[f"{col_letter}{r}"]
            cell.fill = header_fill if r == 1 else header_fill_accent
            cell.font = header_font if r == 1 else header_font_sub
            cell.alignment = header_align
            cell.border = thin_border

    # Merge vertical and horizontal header sections
    ws.merge_cells('A1:A2')
    ws.merge_cells('B1:B2')
    ws.merge_cells('C1:C2')
    ws.merge_cells('D1:D2')
    ws.merge_cells('E1:G1')
    ws.merge_cells('H1:H2')
    ws.merge_cells('I1:I2')
    ws.merge_cells('J1:J2')
    ws.merge_cells('K1:K2')

    data_font = Font(name="Calibri", size=10)
    data_border = Border(
        left=Side(style='thin', color=EXCEL_STYLES['DATA_BORDER']),
        right=Side(style='thin', color=EXCEL_STYLES['DATA_BORDER']),
        top=Side(style='thin', color=EXCEL_STYLES['DATA_BORDER']),
        bottom=Side(style='thin', color=EXCEL_STYLES['DATA_BORDER'])
    )
    zebra_fill = PatternFill(
        start_color=EXCEL_STYLES['ZEBRA_FILL'],
        end_color=EXCEL_STYLES['ZEBRA_FILL'],
        fill_type="solid"
    )

    for r_idx, row_values in enumerate(df_profit.itertuples(index=False), start=3):
        ws.row_dimensions[r_idx].height = 20
        is_even = (r_idx % 2 == 0)

        for c_idx, val in enumerate(row_values, start=1):
            clean_val = val
            if pd.isna(val) or val is None:
                clean_val = ""

            cell = ws.cell(row=r_idx, column=c_idx, value=clean_val)
            cell.font = data_font
            cell.border = data_border
            if is_even:
                cell.fill = zebra_fill

            # Format Currency columns (3: Payout, 9: Cost, 10: Settlement, 11: Profit)
            if c_idx in [3, 9, 10, 11]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if clean_val != "":
                    cell.number_format = '#,##0.00'
            # Format Units columns (4: Gross, 5: RTO, 6: RVP, 7: Cancelled, 8: Net)
            elif c_idx in [4, 5, 6, 7, 8]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if clean_val != "":
                    cell.number_format = '#,##0'
            elif c_idx == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Column Widths
    ws.column_dimensions['A'].width = 38
    ws.column_dimensions['B'].width = 24
    ws.column_dimensions['C'].width = 16
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 18
    ws.column_dimensions['G'].width = 16
    ws.column_dimensions['H'].width = 14
    ws.column_dimensions['I'].width = 24
    ws.column_dimensions['J'].width = 26
    ws.column_dimensions['K'].width = 28


def style_sku_summary_sheet(ws, df_summary):
    """
    Applies standard 1-row header styling to the 'SKU Summary' sheet.
    """
    ws.freeze_panes = "B2"

    header_fill = PatternFill(
        start_color=EXCEL_STYLES['HEADER_FILL_PRIMARY'],
        end_color=EXCEL_STYLES['HEADER_FILL_PRIMARY'],
        fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color=EXCEL_STYLES['HEADER_FONT_COLOR'])
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = list(df_summary.columns)
    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

    data_font = Font(name="Calibri", size=10)
    data_border = Border(
        left=Side(style='thin', color=EXCEL_STYLES['DATA_BORDER']),
        right=Side(style='thin', color=EXCEL_STYLES['DATA_BORDER']),
        top=Side(style='thin', color=EXCEL_STYLES['DATA_BORDER']),
        bottom=Side(style='thin', color=EXCEL_STYLES['DATA_BORDER'])
    )
    zebra_fill = PatternFill(
        start_color=EXCEL_STYLES['ZEBRA_FILL'],
        end_color=EXCEL_STYLES['ZEBRA_FILL'],
        fill_type="solid"
    )

    for r_idx, row_values in enumerate(df_summary.itertuples(index=False), start=2):
        ws.row_dimensions[r_idx].height = 20
        is_even = (r_idx % 2 == 0)

        for c_idx, val in enumerate(row_values, start=1):
            clean_val = val
            if pd.isna(val) or val is None:
                clean_val = ""

            cell = ws.cell(row=r_idx, column=c_idx, value=clean_val)
            cell.font = data_font
            cell.border = data_border
            if is_even:
                cell.fill = zebra_fill

            col_name = headers[c_idx - 1]
            if col_name in ['Bank Settlement', 'PAYOUT']:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if clean_val != "":
                    cell.number_format = '#,##0.00'
            elif col_name in ['Gross Units', 'RTO (Logistics Return)', 'RVP (Customer Return)', 'Cancelled Units', 'Net Units']:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if clean_val != "":
                    cell.number_format = '#,##0'
            elif col_name == 'STYLE NAME':
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col_idx, col_name in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        if col_name == 'SKU Name':
            ws.column_dimensions[col_letter].width = 38
        elif col_name == 'STYLE NAME':
            ws.column_dimensions[col_letter].width = 24
        elif col_name in ['Bank Settlement', 'PAYOUT']:
            ws.column_dimensions[col_letter].width = 18
        else:
            ws.column_dimensions[col_letter].width = 16


def export_to_styled_excel(df_summary, df_profit=None):
    """
    Creates an in-memory multi-sheet Excel workbook buffer containing:
      Sheet 1: 'Profit & Cost Analysis' (2-Row merged header)
      Sheet 2: 'SKU Summary' (1-Row header)

    Returns:
        io.BytesIO: Ready-to-download or save binary buffer.
    """
    wb = openpyxl.Workbook()

    if df_profit is None:
        df_profit = compute_profit_analysis_df(df_summary)

    # Sheet 1: Profit & Cost Analysis
    ws_profit = wb.active
    ws_profit.title = "Profit & Cost Analysis"
    style_profit_analysis_sheet(ws_profit, df_profit)

    # Sheet 2: SKU Summary
    ws_summary = wb.create_sheet(title="SKU Summary")
    style_sku_summary_sheet(ws_summary, df_summary)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def save_styled_excel_file(df_summary, output_path, df_profit=None):
    """Saves multi-sheet styled Excel workbook directly to disk."""
    buffer = export_to_styled_excel(df_summary, df_profit=df_profit)
    with open(output_path, 'wb') as f:
        f.write(buffer.read())
